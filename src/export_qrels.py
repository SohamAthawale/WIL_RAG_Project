"""Export relevance judgements from the annotation workbook into the test set.

The metrics read `data/testset/test_set.json`, not the spreadsheet. Until the
judgements land there, a fully judged workbook unblocks nobody.

Writes three fields per question:

  gold_chunk_ids     chunk ids graded 1 or higher — for Recall@k and MRR
  gold_chunk_grades  chunk id -> grade 0..3       — for NDCG, which needs the scale
  wrong_subclass_chunks  chunks graded 0 *because* they belong to another visa,
                     as distinct from merely off-topic. Feeds the confusion work.

Provenance is recorded alongside: judgements name specific chunks in a specific
corpus, so a snapshot change invalidates them. Recording the snapshot makes that
detectable rather than silent.

Usage:
    python export_qrels.py                      # from the annotated workbook
    python export_qrels.py --sheet <path>
    python export_qrels.py --check              # report staleness, write nothing
"""

import argparse
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
TESTSET = ROOT / "data" / "testset" / "test_set.json"
MANIFEST = ROOT / "data" / "documents" / "MANIFEST.json"
DEFAULT_SHEET = ROOT / "results" / "annotation_sheet_annotated.xlsx"


def read_judgements(sheet: Path) -> tuple[dict, dict]:
    wb = load_workbook(sheet)
    chunks, answers = wb["Chunks"], wb["Answers"]

    grades: dict[str, dict[str, int]] = defaultdict(dict)
    wrong_sc: dict[str, list[str]] = defaultdict(list)
    annotators: set[str] = set()
    ungraded = 0

    for row in chunks.iter_rows(min_row=2, values_only=True):
        qid, _q, chunk_id, _src, _txt, _score, _top3, grade, wrong, annotator = row[:10]
        if grade in (None, ""):
            ungraded += 1
            continue
        grades[qid][chunk_id] = int(grade)
        if str(wrong).strip().lower() == "yes":
            wrong_sc[qid].append(chunk_id)
        if annotator:
            annotators.add(str(annotator).strip())

    answer_grades = {}
    for row in answers.iter_rows(min_row=2, values_only=True):
        qid = row[0]
        if row[8] in (None, "") and row[9] in (None, ""):
            continue
        answer_grades[qid] = {
            "b0_correct": row[8], "rag_correct": row[9], "answerability": row[10],
            "wrong_subclass_used": row[11], "citation_grounded": row[12], "annotator": row[13],
        }

    meta = {
        "annotators": sorted(annotators),
        "ungraded_chunk_rows": ungraded,
        "questions_with_judgements": len(grades),
        "answers_graded": len(answer_grades),
    }
    return {"grades": dict(grades), "wrong_subclass": dict(wrong_sc),
            "answers": answer_grades}, meta


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", type=Path, default=DEFAULT_SHEET)
    ap.add_argument("--check", action="store_true", help="report status, write nothing")
    args = ap.parse_args()

    if not args.sheet.exists():
        raise SystemExit(f"No workbook at {args.sheet}")

    snapshot = json.loads(MANIFEST.read_text())["snapshot_id"]
    testset = json.loads(TESTSET.read_text())
    j, meta = read_judgements(args.sheet)

    print(f"Workbook: {args.sheet.name}")
    print(f"  {meta['questions_with_judgements']} questions judged, "
          f"{meta['answers_graded']} answers graded")
    print(f"  annotators: {', '.join(meta['annotators']) or 'NONE RECORDED'}")
    if meta["ungraded_chunk_rows"]:
        print(f"  {meta['ungraded_chunk_rows']} chunk rows still ungraded")

    existing = sum(1 for t in testset if t.get("gold_chunk_ids"))
    if existing:
        prior = {t.get("qrels_snapshot") for t in testset if t.get("qrels_snapshot")}
        print(f"  {existing} questions already carry judgements "
              f"(snapshot {', '.join(prior) or 'unrecorded'})")

    if args.check:
        stale = [t["id"] for t in testset
                 if t.get("gold_chunk_ids") and t.get("qrels_snapshot") not in (None, snapshot)]
        print(f"\nCorpus snapshot: {snapshot}")
        print(f"Stale judgements: {len(stale)}" + (f" — {stale}" if stale else ""))
        return

    updated = 0
    for t in testset:
        g = j["grades"].get(t["id"])
        if not g:
            continue
        t["gold_chunk_ids"] = sorted(cid for cid, grade in g.items() if grade >= 1)
        t["gold_chunk_grades"] = {cid: grade for cid, grade in sorted(g.items())}
        t["wrong_subclass_chunks"] = sorted(j["wrong_subclass"].get(t["id"], []))
        # Judgements name chunks in a specific corpus. Without this, a snapshot
        # change silently invalidates them - the failure this project keeps hitting.
        t["qrels_snapshot"] = snapshot
        t["qrels_recorded_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        updated += 1

    TESTSET.write_text(json.dumps(testset, indent=2) + "\n")

    total_rel = sum(len(t.get("gold_chunk_ids", [])) for t in testset)
    print(f"\nWrote judgements for {updated} questions to {TESTSET.relative_to(ROOT)}")
    print(f"  {total_rel} chunks graded 1+ across the set, snapshot {snapshot}")

    ans_out = ROOT / "results" / "answer_grades.json"
    ans_out.write_text(json.dumps(
        {"snapshot_id": snapshot, "annotators": meta["annotators"], "grades": j["answers"]},
        indent=2) + "\n")
    print(f"  answer grades -> {ans_out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
