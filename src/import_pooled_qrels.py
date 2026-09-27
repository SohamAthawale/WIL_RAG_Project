"""Import the pooled judging sheet into the test set as qrels.

The 12 questions judged in T02 were judged exhaustively -- every question against all 24
chunks -- so a chunk absent from their grades means the annotator did not see it, which
never happens. The 88 judged here were judged over a pool, so a chunk absent from their
grades means no retriever surfaced it in its top 5 and it is *assumed* not relevant.

Those are different claims and metrics must not silently mix them, so every question now
carries qrels_method. Recall computed over pooled qrels is an upper bound: a relevant
chunk no retriever surfaced was never judged and so cannot count against recall.

The exhaustive rows are left untouched. Re-judging them would void five results already
built on them.
"""
from __future__ import annotations

import json
import re
import shutil
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TESTSET = ROOT / "data" / "testset" / "test_set.json"
MANIFEST = ROOT / "data" / "documents" / "MANIFEST.json"
SHEET = Path("/Users/sohamathawale/Desktop/pooled_judging_sheet_assigned.xlsx")


def subclass_of(name: str) -> str | None:
    m = re.search(r"subclass_(\d+)", name or "")
    return m.group(1) if m else None


def main() -> None:
    snapshot = json.loads(MANIFEST.read_text())["snapshot_id"]
    testset = json.loads(TESTSET.read_text())
    by_id = {t["id"]: t for t in testset}

    ws = openpyxl.load_workbook(SHEET)["Pool"]
    hdr = [c.value for c in ws[1]]
    i = {h: n for n, h in enumerate(hdr)}

    grades: dict[str, dict[str, int]] = defaultdict(dict)
    annotators: dict[str, set[str]] = defaultdict(set)
    for r in ws.iter_rows(min_row=2, values_only=True):
        qid, cid = r[i["question_id"]], r[i["chunk_id"]]
        grades[qid][cid] = int(r[i["relevance"]])
        annotators[qid].add(str(r[i["annotator"]]).strip())

    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    shutil.copy(TESTSET, TESTSET.with_suffix(".json.bak"))

    touched = skipped = 0
    for qid, g in grades.items():
        t = by_id.get(qid)
        if t is None:
            continue
        if t.get("gold_chunk_ids"):
            # Already judged exhaustively in T02. Leave it alone.
            skipped += 1
            continue
        gold_source = t.get("gold_source", "")
        gold_sc = subclass_of(gold_source)
        t["gold_chunk_ids"] = sorted(c for c, v in g.items() if v >= 1)
        t["gold_chunk_grades"] = {c: v for c, v in sorted(g.items())}
        t["wrong_subclass_chunks"] = sorted(
            c for c, v in g.items()
            if v == 0 and gold_sc and subclass_of(c) and subclass_of(c) != gold_sc
        )
        t["qrels_snapshot"] = snapshot
        t["qrels_recorded_at"] = now
        t["qrels_method"] = "pooled_top5_union_dense_bm25_rrf"
        t["qrels_annotators"] = sorted(annotators[qid])
        touched += 1

    # Mark the earlier rows so the two methods are never confused downstream.
    for t in testset:
        if t.get("gold_chunk_ids") and "qrels_method" not in t:
            t["qrels_method"] = "exhaustive_all_chunks"

    TESTSET.write_text(json.dumps(testset, indent=2) + "\n")

    judged = [t for t in testset if t.get("gold_chunk_ids")]
    pooled = [t for t in judged if t["qrels_method"].startswith("pooled")]
    exh = [t for t in judged if t["qrels_method"].startswith("exhaustive")]
    print(f"imported {touched} pooled questions, left {skipped} exhaustive ones untouched")
    print(f"judged now: {len(judged)} of {len(testset)}  "
          f"({len(exh)} exhaustive, {len(pooled)} pooled)")
    print(f"graded chunk rows: {sum(len(t['gold_chunk_grades']) for t in judged)}")
    print(f"backup: {TESTSET.with_suffix('.json.bak').name}")


if __name__ == "__main__":
    main()
